import os 
import json
import shutil
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client = Groq(api_key=my_api_key)
model = "llama-3.3-70b-versatile"

# ---------------------------------------------------------
# 1. Pydantic Schemas
# ---------------------------------------------------------
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []

resume_schema = Resume.model_json_schema()

# ---------------------------------------------------------
# 2. File Readers & Resume Parser
# ---------------------------------------------------------
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

def read_resume(file_path):
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(file_path)
    elif suffix == ".docx":
        return read_docx(file_path)
    elif suffix == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    else:
        return None

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.
    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Return ONLY valid JSON matching this schema:
    {resume_schema}

    Important rules:
    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"Parse the following resume:\n\n{resume_text}"
    
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt}
    ]
    
    response = client.chat.completions.create(
        model=model, 
        messages=messages, 
        response_format={"type": "json_object"}
    )
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    return Resume(**data)

# ---------------------------------------------------------
# 3. FastAPI App & In-Memory Storage
# ---------------------------------------------------------
app = FastAPI(title="AI Portfolio Chatbot")

# CORS middleware to allow connection from Frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Global memory to store current profile & chat history
current_candidate_profile: Resume | None = None
chat_history: list[dict] = []

class MessagePayload(BaseModel):
    message: str

# ---------------------------------------------------------
# 4. Endpoints
# ---------------------------------------------------------

@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    global current_candidate_profile, chat_history
    file_path = Path(f"temp_{file.filename}")
    
    try:
        # Save temp file
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        extracted_text = read_resume(file_path)
        if not extracted_text:
            raise HTTPException(status_code=400, detail="Unsupported file type.")
            
        # Parse resume and store in memory
        current_candidate_profile = parse_resume(extracted_text)
        chat_history = [] # Reset memory on new upload
        
        return {
            "status": "success",
            "message": "Resume uploaded and parsed successfully!",
            "parsed_profile": current_candidate_profile.model_dump()
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if file_path.exists():
            file_path.unlink() # Cleanup local temp file

@app.post("/chat")
def chat_with_ai(payload: MessagePayload):
    global current_candidate_profile, chat_history
    
    if not current_candidate_profile:
        context_str = "No resume uploaded yet. Tell the user to upload a resume first."
    else:
        context_str = current_candidate_profile.model_dump_json()

    # System prompt for portfolio AI agent
    system_prompt = f"""
    You are an AI representative for the candidate. Your goal is to represent the candidate to HRs/recruiters professionally and accurately.

    Candidate Profile:
    {context_str}

    Rules:
    1. Be polite, concise, and executive in your answers.
    2. Answer questions STRICTLY based on the provided profile. Do NOT fabricate skills, experiences, or project details.
    3. If the answer is not present in the profile, say: "I don't have that detail in my candidate file currently."
    4. Highlight key strengths when relevant.
    """

    # Build prompt with history
    messages = [{"role": "system", "content": system_prompt}]
    messages.extend(chat_history)
    messages.append({"role": "user", "content": payload.message})

    # Call LLM
    response = client.chat.completions.create(model=model, messages=messages)
    reply = response.choices[0].message.content

    # Update conversation history
    chat_history.append({"role": "user", "content": payload.message})
    chat_history.append({"role": "assistant", "content": reply})

    return {"reply": reply}
@app.get("/")
def home():
    return {"message": "AI Portfolio API is running! Go to /docs to test endpoints."}